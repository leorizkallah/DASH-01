"""Build the controller midterm deck, embedded clips, figures and source notes.

Requires python-pptx, matplotlib, Pillow, imageio and imageio-ffmpeg.
Run from the workspace with .venv/Scripts/python report/scripts/build_midterm_controller.py.
"""

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess

import imageio.v2 as imageio
import imageio_ffmpeg
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "report/presentations/midterm-2026-10-29"
ASSETS = OUT / "assets"
DECK_STEM = "DASH01_controller_midterm_Tim_template"
BG, DARK, RED, GREY = "F7F8F5", "26323F", "E2001A", "5A626D"
PALE, BLUE, ORANGE, GREEN = "E8ECEA", "168BD2", "EF8018", "176A55"
FONT = "Arial"
SECTIONS = ["Context", "What was done", "Objectives and timeline", "Challenges"]
SLIDE_NOTES = []
INPUTS = set()
TEMPLATE_SLIDES = {}
BODY_START = {}
LAST_SECTION = None
OUTLINE_ADDED = False
CLOSING_ADDED = False


def replace_template_text(shape, value):
    """Keep Tim's run/paragraph formatting while changing the words."""
    frame = shape.text_frame
    p = frame.paragraphs[0]
    if not p.runs:
        p.text = value
    else:
        p.runs[0].text = value
        for run in list(p.runs)[1:]:
            p._p.remove(run._r)
    for paragraph in list(frame.paragraphs)[1:]:
        frame._txBody.remove(paragraph._p)


def template_clone(prs, name):
    """Clone a supplied slide with its original layout and internal media links."""
    src = TEMPLATE_SLIDES[name]
    dst = prs.slides.add_slide(src.slide_layout)
    for shape in list(dst.shapes):
        dst.shapes._spTree.remove(shape._element)
    links = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        target = rel.target_ref if rel.is_external else rel.target_part
        links[rel.rId] = dst.part.relate_to(
            target, rel.reltype, is_external=rel.is_external
        )
    for shape in src.shapes:
        node = deepcopy(shape._element)
        for element in node.iter():
            for key, value in list(element.attrib.items()):
                if (
                    key.startswith(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
                    )
                    and value in links
                ):
                    element.set(key, links[value])
        dst.shapes._spTree.insert_element_before(node, "p:extLst")
    for shape in dst.shapes:
        if (
            shape.has_text_frame
            and shape.text.strip().isdigit()
            and shape.top > Inches(6.5)
        ):
            replace_template_text(shape, str(len(prs.slides)))
    return dst


def add_outline(prs):
    s = template_clone(prs, "outline")
    replacements = {
        "The robot · Semester goal · Starting point": "Existing controller · Semester goal · Starting point",
        "Choices · Mechanics · Foot · Safety · Simulation": "Controller · Learning · Gait · Results",
        "Second half of the semester": "Demo sequence · Early treadmill tests",
        "Open problems · Support needed": "Sim-to-real gap · Unknown unknowns",
    }
    for shape in s.shapes:
        if shape.has_text_frame and shape.text in replacements:
            replace_template_text(shape, replacements[shape.text])
    note(
        s,
        10,
        "I will follow the same four parts as Tim: context, what was done, objectives and timeline, and challenges. My focus is the controller and its validation.",
        [
            "Outline cloned from Tim Kobler’s supplied presentation; controller-specific subtitles."
        ],
    )


def add_divider(prs, section):
    s = template_clone(prs, f"part{section + 1}")
    subtitles = [
        "Existing controller · Semester goal · Starting point",
        "Model · Learning · Gait · Results",
        "Demo sequence · Early treadmill tests",
        "Sim-to-real gap · Unknown unknowns",
    ]
    for shape in s.shapes:
        if shape.has_text_frame and abs(shape.top / Inches(1) - 4.097) < 0.01:
            replace_template_text(shape, subtitles[section])
    note(
        s,
        5,
        f"Part {section + 1}: {SECTIONS[section]}.",
        [
            "Section divider cloned from the corresponding part of Tim’s supplied deck, including navigation, layout and typography."
        ],
    )


def add_closing(prs):
    s = template_clone(prs, "closing")
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith("Next:"):
            replace_template_text(shape, "Next: treadmill tests and sim-to-real")
    note(
        s,
        5,
        "Thank you. I am happy to discuss the controller, the experimental results and the first treadmill tests.",
        [
            "Closing slide cloned from Tim’s supplied presentation; controller-specific subtitle."
        ],
    )


def fit_content_to_template(prs):
    """Fit existing editable diagrams/media into Tim's light-slide content area."""
    for slide in prs.slides:
        start = BODY_START.get(slide.slide_id)
        if start is None:
            continue
        for shape in list(slide.shapes)[start:]:
            x, y, w, h = [
                getattr(shape, k) / Inches(1)
                for k in ("left", "top", "width", "height")
            ]
            shape.left = Inches(0.889 + (x - 0.65) * (11.556 / 12.02))
            shape.top = Inches(1.80 + (y - 1.80) * 0.94)
            shape.width = Inches(w * (11.556 / 12.02))
            shape.height = Inches(h * 0.94)
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.font.size:
                        paragraph.font.size = Pt(paragraph.font.size.pt * 0.94)
                    if paragraph.space_after:
                        paragraph.space_after = Pt(paragraph.space_after.pt * 0.65)
                    for run in paragraph.runs:
                        run.font.name = FONT
                        if run.font.size:
                            run.font.size = Pt(run.font.size.pt * 0.94)


def read(relative):
    p = ROOT / relative
    INPUTS.add(p)
    return json.loads(p.read_text(encoding="utf-8"))


def rgb(color):
    return RGBColor.from_string(color)


def rect(slide, x, y, w, h, color=PALE, radius=False, line=None):
    s = slide.shapes.add_shape(
        SHAPE.ROUNDED_RECTANGLE if radius else SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    if line:
        s.line.color.rgb = rgb(line)
    else:
        s.line.fill.background()
    if radius:
        s.adjustments[0] = 0.08
    return s


def text(slide, value, x, y, w, h, size=22, bold=False, color=DARK, align=None):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size = FONT, Pt(size)
        p.font.bold, p.font.color.rgb = bold, rgb(color)
        p.space_after = Pt(7)
        if align is not None:
            p.alignment = align
    return s


def picture(slide, path, x, y, w, h):
    path = Path(path)
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    ww, hh = iw * scale, ih * scale
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - ww) / 2),
        Inches(y + (h - hh) / 2),
        width=Inches(ww),
        height=Inches(hh),
    )


def arrow(slide, x, y, w=0.35, h=0.20, color=GREY):
    return rect_arrow(slide, x, y, w, h, color)


def rect_arrow(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(
        SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    return s


def note(slide, timing, script, sources):
    block = (
        f"Suggested time: {timing} seconds.\n\n{script}\n\nEvidence / provenance:\n"
        + "\n".join(sources)
    )
    notes = slide.notes_slide
    if notes.notes_text_frame is None:
        # Tim's supplied notes master has no body placeholder. Add one while
        # preserving that master, so the controller speaker notes remain editable.
        shape_id = max([sh.shape_id for sh in notes.shapes] + [1]) + 1
        body = parse_xml(
            f'<p:sp {nsdecls("p", "a")}><p:nvSpPr><p:cNvPr id="{shape_id}" name="Speaker Notes"/><p:cNvSpPr/><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
        )
        notes.shapes._spTree.insert_element_before(body, "p:extLst")
    notes.notes_text_frame.text = block
    SLIDE_NOTES.append(
        dict(
            number=len(SLIDE_NOTES) + 1, seconds=timing, script=script, sources=sources
        )
    )


def base(prs, title, section, takeaway=None, backup=False):
    global OUTLINE_ADDED, LAST_SECTION, CLOSING_ADDED
    if not backup:
        if not OUTLINE_ADDED:
            add_outline(prs)
            OUTLINE_ADDED = True
        if section != LAST_SECTION:
            add_divider(prs, section)
            LAST_SECTION = section
    elif not CLOSING_ADDED:
        add_closing(prs)
        CLOSING_ADDED = True
    titles = {
        "One robot, one shared goal: cross the sim-to-real gap": "Semester goal",
        "Nathann’s controller: learned gait + fast corrections": "Starting point: Nathann’s controller",
        "Experiments: from “stays alive” to useful walking": "What we explored and what we learned",
        "The evolution is visible — but videos need measurements": "Walking policies: visual evolution",
        "Speed tracking improved substantially": "Speed tracking",
        "Foot lift is learnable on this robot": "Foot clearance: first results",
        "Next: lower, more consistent clearance": "Next training objective",
        "Robotics Day: the original three-phase demo": "Objectives for Robotics Day",
        "Start treadmill tests early — learn in parallel": "Timeline",
        "The main challenge: the sim-to-real gap": "Where we are stuck",
        "A stronger simulation baseline. Now test the real robot.": "What we need",
        "Backup: what the rewards actually measure": "Backup: reward definitions",
        "Backup: keep the failed replay visible": "Backup: the failed replay",
        "Backup: evidence and comparison limits": "Backup: evidence and comparison limits",
    }
    s = prs.slides.add_slide(next(l for l in prs.slide_layouts if l.name == "LIGHT"))
    source_title = next(
        sh
        for sh in TEMPLATE_SLIDES["light"].shapes
        if sh.has_text_frame and sh.text == "What is DASH-01?"
    )
    s.shapes._spTree.insert_element_before(deepcopy(source_title._element), "p:extLst")
    replace_template_text(s.shapes[-1], titles.get(title, title))
    text(
        s,
        str(len(prs.slides)),
        10.361,
        6.597,
        2.083,
        0.417,
        20,
        color=GREY,
        align=PP_ALIGN.RIGHT,
    )
    if takeaway:
        text(s, takeaway, 0.889, 6.22, 11.556, 0.31, 16, True, color=GREY)
    if backup:
        s._element.set("show", "0")
    BODY_START[s.slide_id] = len(s.shapes)
    return s


def card(slide, x, y, w, h, kicker, heading, body, accent=RED):
    rect(slide, x, y, w, h, "FFFFFF", True)
    rect(slide, x, y, 0.055, h, accent)
    text(slide, kicker, x + 0.20, y + 0.17, w - 0.40, 0.28, 12, True, accent)
    text(slide, heading, x + 0.20, y + 0.61, w - 0.40, 0.7, 25, True)
    text(slide, body, x + 0.20, y + 1.42, w - 0.40, h - 1.55, 20)


def make_clip(name, source, start, duration, crop=True):
    src = ROOT / source
    INPUTS.add(src)
    dst = ASSETS / f"{name}.mp4"
    filters = "crop=1280:596:0:80,scale=1152:-2" if crop else "scale=1152:-2"
    if not dst.exists():
        subprocess.run(
            [
                imageio_ffmpeg.get_ffmpeg_exe(),
                "-v",
                "error",
                "-nostdin",
                "-y",
                "-ss",
                str(start),
                "-i",
                str(src),
                "-t",
                str(duration),
                "-an",
                "-vf",
                filters,
                "-r",
                "20",
                "-c:v",
                "libx264",
                "-crf",
                "25",
                "-preset",
                "fast",
                "-threads",
                "4",
                "-pix_fmt",
                "yuv420p",
                "-movflags",
                "+faststart",
                str(dst),
            ],
            check=True,
        )
    poster = ASSETS / f"{name}.jpg"
    with imageio.get_reader(dst) as r:
        info = r.get_meta_data()
        Image.fromarray(
            r.get_data(min(20, round(info["duration"] * info["fps"]) - 1))
        ).save(poster, quality=92)
    return dst, poster


def movie(slide, clip, x, y, w, h):
    return slide.shapes.add_movie(
        str(clip[0]),
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        poster_frame_image=str(clip[1]),
        mime_type="video/mp4",
    )


def figures():
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 14,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.facecolor": "#" + BG,
            "figure.facecolor": "#" + BG,
            "axes.labelcolor": "#" + DARK,
            "text.color": "#" + DARK,
            "xtick.color": "#" + GREY,
            "ytick.color": "#" + GREY,
        }
    )
    old = read("docs/experiments/walk01_20261019_v1/best_eval.json")
    flat = read("docs/experiments/dash01_flatfeet_symmetry_v1/best_eval.json")
    fig, ax = plt.subplots(figsize=(9.2, 4.6))
    for r, label, col in [(old, "WALK-01", GREY), (flat, "Flat-feet baseline", RED)]:
        a = np.asarray(r["per_cmd"])
        ax.plot(a[:, 0], a[:, 1], "-o", color="#" + col, lw=3, ms=8, label=label)
    ax.set(
        xlabel="Commanded speed (m/s)",
        ylabel="Mean absolute speed error (m/s)",
        ylim=(0, 0.46),
        xticks=[0.25, 0.5, 0.75, 1, 1.5, 2, 2.5],
    )
    ax.grid(axis="y", alpha=0.18)
    ax.legend(frameon=False, loc="upper left")
    fig.tight_layout()
    fig.savefig(ASSETS / "tracking.png", dpi=180)
    fig.savefig(ASSETS / "tracking.svg")
    plt.close(fig)
    evpath = (
        ROOT / "docs/experiments/clearance01_20261025_v1/clearance_evaluations.jsonl"
    )
    INPUTS.add(evpath)
    ev = [json.loads(s) for s in evpath.read_text().splitlines()]
    x = np.array([r["step"] for r in ev]) / 1e6
    fig, axes = plt.subplots(2, 1, figsize=(8.7, 5.1), sharex=True)
    for i, side, col in [(0, "Left", BLUE), (1, "Right", ORANGE)]:
        axes[0].plot(
            x,
            [r["mean_air_clearance_mm"][i] for r in ev],
            "-o",
            color="#" + col,
            lw=2.6,
            ms=4,
            label=side,
        )
        axes[1].plot(
            x,
            [100 * r["midswing_ge20mm_fraction"][i] for r in ev],
            "-o",
            color="#" + col,
            lw=2.6,
            ms=4,
        )
    axes[0].axhline(20, color="#" + GREY, ls="--")
    axes[0].set(ylabel="Mean lift (mm)", ylim=(0, 28))
    axes[0].legend(loc="upper left", frameon=False, ncol=2)
    axes[1].axhline(80, color="#" + GREY, ls="--")
    axes[1].text(1, 84, "80% criterion", fontsize=11, color="#" + GREY)
    axes[1].set(
        ylabel="Mid-swing ≥20 mm (%)",
        xlabel="Additional training steps (millions)",
        ylim=(0, 100),
    )
    for ax in axes:
        ax.grid(alpha=0.16)
    fig.tight_layout()
    fig.savefig(ASSETS / "clearance.png", dpi=180)
    fig.savefig(ASSETS / "clearance.svg")
    plt.close(fig)


def build(reference):
    global LAST_SECTION, OUTLINE_ADDED, CLOSING_ADDED
    # Support repeated builds from an imported module as well as the CLI.
    LAST_SECTION = None
    OUTLINE_ADDED = CLOSING_ADDED = False
    SLIDE_NOTES.clear()
    TEMPLATE_SLIDES.clear()
    BODY_START.clear()
    INPUTS.clear()
    ASSETS.mkdir(parents=True, exist_ok=True)
    if not (ASSETS / "epfl.png").exists():
        ref = Presentation(reference)
        for sh in ref.slides[0].shapes:
            if sh.name in ("Image 0", "Image 1"):
                (
                    ASSETS
                    / ("epfl.png" if sh.name == "Image 0" else "robot_reference.png")
                ).write_bytes(sh.image.blob)
    figures()
    walk = make_clip("walk01", "report/videos/walk01_best_1.50mps.mp4", 6, 12)
    flat = make_clip("flatfeet", "report/videos/flatfeet_best_1.50mps.mp4", 6, 12)
    lifted = make_clip(
        "clearance01", "report/videos/clearance01_final_1.50mps.mp4", 6, 12
    )
    failed = make_clip(
        "standalone_failure", "report/videos/clearance01_standalone_1.50mps.mp4", 0, 3.7
    )
    # Reuse the actual supplied theme, masters and layouts.
    prs = Presentation(reference)
    TEMPLATE_SLIDES.update(
        {
            name: prs.slides[i]
            for name, i in {
                "cover": 0,
                "outline": 1,
                "part1": 2,
                "light": 3,
                "part2": 6,
                "part3": 12,
                "part4": 15,
                "closing": 18,
            }.items()
        }
    )
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "DASH-01 — Controller midterm"
    prs.core_properties.author = "Léo-Paul Rizkallah"
    prs.core_properties.subject = (
        "Controller architecture, experiments, demo and sim-to-real roadmap"
    )
    prs.core_properties.comments = "Prepared 26 October 2026. Results are simulation-only unless explicitly stated. Videos are embedded."

    # Cover cloned from the supplied presentation.
    s = template_clone(prs, "cover")
    for shape in s.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text == "Tim Kobler":
            replace_template_text(shape, "L" + chr(233) + "o-Paul Rizkallah")
        elif shape.text == "Midterm Presentation":
            replace_template_text(
                shape, "Midterm Presentation " + chr(183) + " Controller"
            )
    text(
        s,
        "Results reviewed 26 October 2026",
        0.889,
        6.40,
        6.944,
        0.28,
        14,
        color="AEB6BF",
    )
    text(
        s,
        "Previous robot iteration - photo from Tim",
        8.562,
        6.82,
        3.847,
        0.25,
        11,
        color="AEB6BF",
    )
    note(
        s,
        20,
        "My project continues Nathann’s controller work on DASH-01, alongside Tim’s mechanical redesign. The first part of the semester has been about adapting the simulator and recovering useful locomotion on the new robot. I will show the architecture we inherited, what the experiments taught us, and how we move from promising simulation videos to treadmill tests and a short demonstration.",
        [
            "Reference style, lab/supervision and photograph: Tim Kobler, DASH01_midterm.pptx, supplied by user.",
            "Data cutoff: 26 October 2026; presentation date follows Tim’s deck.",
        ],
    )

    # 2 — context / goal
    s = base(
        prs,
        "One robot, one shared goal: cross the sim-to-real gap",
        0,
        "My focus: adapt and validate the learned controller. Tim’s focus: the mechanics.",
    )
    card(
        s,
        0.65,
        1.95,
        3.87,
        3.90,
        "INHERITED",
        "Working baseline",
        "Nathann’s RL pipeline\nGait generator + feedback\nDeployment code and UI",
    )
    card(
        s,
        4.73,
        1.95,
        3.87,
        3.90,
        "WHAT CHANGED",
        "A new model",
        "New geometry and flat feet\nCorrected mass and inertia\nNew homing configuration",
    )
    card(
        s,
        8.81,
        1.95,
        3.87,
        3.90,
        "SEMESTER OBJECTIVE",
        "A short robot demo",
        "Start upright\nWalk at commanded speeds\nStop and remain standing",
    )
    note(
        s,
        35,
        "The starting point was not an empty software project: Nathann already had a working learning pipeline and a structured walking controller. But new geometry, feet and mass distribution change the dynamics. My part is to adapt the model, train useful behavior, diagnose what the policy actually does and validate it on the real robot. Tim and I therefore share the sim-to-real objective. The final goal remains startup, commanded walking and a controlled stop, rather than just a fast simulation clip.",
        [
            "README.md: original three-phase demonstration plan.",
            "report/cad-v2-import.md: 14.594154 kg corrected simulation model.",
            "report/architecture.md; RLframework/networks.py.",
        ],
    )

    # 3 — inherited architecture, native editable shapes
    s = base(
        prs,
        "Nathann’s controller: learned gait + fast corrections",
        0,
        "The policy generates motion targets; motor feedback turns those targets into torque.",
    )
    blocks = [
        ("Sensors", "IMU + joints\n180 ms history"),
        ("State estimate", "Learned velocity\nFrom sensor history"),
        ("RL policy", "384 obs + 3 est.\n47 outputs · 100 Hz"),
        ("Gait generator", "41 gait parameters\n6 joint corrections"),
        ("PD + limits", "Targets + gains\nTorque/speed caps"),
    ]
    for i, (a, b) in enumerate(blocks):
        x = 0.65 + i * 2.44
        rect(s, x, 2.45, 2.23, 1.45, "FFFFFF", True)
        text(s, a, x + 0.13, 2.64, 1.97, 0.4, 20, True, RED if i == 2 else DARK)
        text(s, b, x + 0.13, 3.12, 1.97, 0.70, 16)
        if i < 4:
            arrow(s, x + 2.25, 3.02, 0.17, 0.19)
    rect(s, 5.53, 1.84, 2.23, 0.40, PALE, True)
    text(s, "Commanded speed", 5.65, 1.91, 2.0, 0.24, 15, True)
    rect(s, 10.41, 4.48, 2.23, 0.95, DARK, True)
    text(s, "Six motors\n+ robot", 10.54, 4.63, 1.97, 0.73, 18, True, "FFFFFF")
    down = s.shapes.add_shape(
        SHAPE.DOWN_ARROW, Inches(11.39), Inches(4.02), Inches(0.23), Inches(0.33)
    )
    down.fill.solid()
    down.fill.fore_color.rgb = rgb(GREY)
    down.line.fill.background()
    rect(s, 0.65, 4.48, 8.95, 0.95, PALE, True)
    text(
        s,
        "Training only: PPO + privileged critic in MuJoCo / MJX",
        0.85,
        4.61,
        8.5,
        0.32,
        21,
        True,
    )
    text(
        s,
        "Ground-truth velocity and contacts help learning, not the deployed actor.",
        0.85,
        5.02,
        8.5,
        0.27,
        17,
        color=GREY,
    )
    text(
        s,
        "Feedback: joint measurements and IMU → next policy update. Simulation advances in 1 ms substeps.",
        0.85,
        5.80,
        11.55,
        0.42,
        18,
        color=GREY,
    )
    note(
        s,
        55,
        "This is a hybrid controller. A neural network uses a short history of joint measurements and the IMU, plus a learned velocity estimate. At one hundred hertz it outputs forty-seven values: forty-one parameters describing periodic joint motion and gains, and six immediate corrective joint offsets. Gait parameters are committed on the gait clock; corrections remain active every tick. PD control and motor limits then generate torque. During PPO training, a separate critic can see simulator ground truth, including contact. The actor cannot. That separation matters for transfer: simulation contact rewards are not additional hardware sensors.",
        [
            "RLframework/networks.py; RLframework/env.py; RLframework/gait.py; RLframework/drive.py.",
            "report/architecture.md: 384 actor observations, 47 actions, estimator 128/64, actor 256/256, privileged critic 413.",
            "Current CAD-v2 policies disable STOP; START/STOP integration remains planned.",
        ],
    )

    # 4 — experiments / history
    s = base(
        prs,
        "Experiments: from “stays alive” to useful walking",
        1,
        "Survival, speed tracking and gait quality are different objectives.",
    )
    entries = [
        (
            "1",
            "Stopping on the previous robot",
            "Abrupt braking exposed a missing standing skill. Imitation improved quiet holding.",
            "Older model",
        ),
        (
            "2",
            "New CAD and gait discovery",
            "Correct inertials and closed loops; investigate dragging, asymmetry and micro-steps.",
            "Model + diagnosis",
        ),
        (
            "3",
            "Reproduce the walking recipe",
            "Recover sustained commanded locomotion on CAD-v2 from the supplied homing pose.",
            "WALK-01",
        ),
        (
            "4",
            "Measure real heel/toe support",
            "Add contact-based slip, completed-cycle symmetry and mid-stance flat support.",
            "Flat-feet baseline",
        ),
        (
            "5",
            "Shape the swing trajectory",
            "Measure clearance; teach higher lift, then target a smaller height band.",
            "CLEARANCE-01 / 02",
        ),
    ]
    for i, (num, a, b, c) in enumerate(entries):
        y = 1.90 + i * 0.84
        rect(s, 0.65, y, 0.46, 0.46, RED, True)
        text(s, num, 0.74, y + 0.065, 0.27, 0.30, 18, True, "FFFFFF")
        text(s, a, 1.35, y - 0.02, 8.7, 0.35, 22, True)
        text(s, b, 1.35, y + 0.36, 10.9, 0.38, 17, color=GREY)
        text(s, c, 10.45, y - 0.01, 2.2, 0.32, 13, True, RED, PP_ALIGN.RIGHT)
    note(
        s,
        55,
        "There were several phases, and the early failures were useful. On the previous robot, simply asking a walking policy to stop did not teach balance. Standing demonstrations improved near-standing holding, but that was not a robust stop from motion. With the new CAD we then saw dragging, unilateral activity and micro-steps, so reward and diagnostic definitions became central. After reproducing sustained walking, we moved to actual heel and toe contact measurements. The latest experiments focus on foot lift. This progression is why I now separate staying upright, following a speed and having an acceptable gait.",
        [
            "docs/stop03-experiment-report.md; docs/stop05-experiment-report.md.",
            "report/cad-v2-import.md; docs/gait11-experiment-report.md; docs/walk01-experiment-report.md.",
            "docs/flatfeet-symmetry-v1.md; docs/clearance01-experiment-report.md; docs/clearance02-experiment-report.md.",
            "STOP-03: 0/32 to 21/32 strict 5 s holds after imitation; near-standing, previous model, not a moving-stop result.",
        ],
    )

    # 5 — two embedded videos
    s = base(
        prs,
        "The evolution is visible — but videos need measurements",
        1,
        "Same 1.50 m/s command; different trained policies. Click each clip to play.",
    )
    text(s, "Walking reproduction", 0.65, 1.88, 5.82, 0.43, 25, True)
    text(s, "WALK-01 best · checkpoint 169.57M", 0.65, 2.34, 5.82, 0.28, 16, color=GREY)
    movie(s, walk, 0.65, 2.91, 5.84, 2.72)
    text(s, "Flat-feet baseline", 6.82, 1.88, 5.82, 0.43, 25, True)
    text(
        s,
        "Best checkpoint 184.32M · later used for fine-tuning",
        6.82,
        2.34,
        5.82,
        0.38,
        16,
        color=GREY,
    )
    movie(s, flat, 6.82, 2.91, 5.84, 2.72)
    text(
        s,
        "Simulation · nominal MJX replays · 12 s excerpts (t = 6–18 s) · 1× playback",
        0.65,
        5.92,
        12,
        0.30,
        16,
        color=GREY,
    )
    note(
        s,
        55,
        "These clips have the same target speed and show two stages of the work. The first is the reproduced walking recipe; the second is the flat-feet baseline. Please look at the feet, rather than only the body moving forward. Even when the policy survives, support can be rolled onto one pad and a foot can slide instead of taking a clean swing. The clips are normal-speed simulation recordings, not hardware experiments. They illustrate the evolution, but they are not a controlled single-reward ablation, and one good clip cannot establish robustness.",
        [
            "report/videos/walk01_best_1.50mps.mp4; report/videos/flatfeet_best_1.50mps.mp4.",
            "docs/experiments/walk01_20261019_v1/; docs/experiments/dash01_flatfeet_symmetry_v1/.",
            "Both source clips: nominal MJX, target 1.50 m/s, seed 1000, first episode retained. Excerpts 6–18 s; header/footer cropped for readability, motion not modified.",
        ],
    )

    # 6 — speed plot
    s = base(
        prs,
        "Speed tracking improved substantially",
        1,
        "A better tracking score does not automatically mean a clean or transferable gait.",
    )
    picture(s, ASSETS / "tracking.png", 0.55, 1.95, 8.65, 4.20)
    rect(s, 9.48, 2.0, 3.18, 3.86, "FFFFFF", True)
    text(s, "Mean speed error", 9.71, 2.25, 2.72, 0.4, 20, True)
    text(s, "0.181 → 0.049", 9.71, 2.85, 2.72, 0.70, 29, True, RED)
    text(s, "m/s across the speed ladder", 9.71, 3.58, 2.72, 0.6, 18, color=GREY)
    text(s, "28 / 28", 9.71, 4.32, 2.72, 0.6, 32, True)
    text(
        s,
        "full 30 s trials\nboth best policies",
        9.71,
        4.98,
        2.72,
        0.7,
        17,
        color=GREY,
    )
    note(
        s,
        40,
        "The numerical comparison is encouraging. On the same nominal seven-speed evaluation ladder, the saved best flat-feet policy reduces average absolute speed error from about zero-point-eighteen to zero-point-zero-five metres per second. Both saved best policies complete all twenty-eight trials, so survival alone would not show this improvement. We must still inspect support, slip and individual leg behavior. Several things changed together in the flat-feet experiment, and these are single-seed training results, so the graph supports progress rather than a causal claim for one isolated reward term.",
        [
            "docs/experiments/walk01_20261019_v1/best_eval.json.",
            "docs/experiments/dash01_flatfeet_symmetry_v1/best_eval.json.",
            "28 nominal trials, 30 s, 5 s excluded from tracking, command ladder .25/.5/.75/1/1.5/2/2.5 m/s. Not a controlled reward ablation.",
        ],
    )

    # 7 — current lift clip
    s = base(
        prs,
        "Foot lift is learnable on this robot",
        1,
        "Promising swing motion in simulation; startup reliability still needs investigation.",
    )
    movie(s, lifted, 0.65, 2.05, 8.12, 3.78)
    text(s, "Mean airborne lift", 9.12, 2.02, 3.55, 0.42, 23, True)
    text(s, "23.6 mm", 9.12, 2.67, 3.55, 0.65, 38, True, BLUE)
    text(s, "left  (baseline: 2.46 mm)", 9.12, 3.34, 3.55, 0.39, 18, color=GREY)
    text(s, "20.9 mm", 9.12, 3.95, 3.55, 0.65, 38, True, ORANGE)
    text(s, "right  (baseline: 2.25 mm)", 9.12, 4.62, 3.55, 0.39, 18, color=GREY)
    text(s, "+50M training steps\nTarget: 1.50 m/s", 9.12, 5.29, 3.55, 0.76, 20, True)
    text(
        s,
        "Evaluation-protocol replay · trial 0 / 12 · t = 6–18 s · 1×",
        0.65,
        6.01,
        8.25,
        0.28,
        15,
        color=GREY,
    )
    note(
        s,
        45,
        "The flat-feet baseline had only two to three millimetres of average airborne clearance. A dense penalty below the target gave a learning signal even when feet were still dragging, and fifty million additional steps increased mean lift to roughly twenty-four and twenty-one millimetres. This clip is the first trial from the retained batch, not a selected best-looking trial. There is an important caveat: a separate standalone replay fell after three-point-seven seconds. We keep that failure in the backup slides. The result demonstrates learnable foot lift, not a validated standalone startup or hardware-ready controller.",
        [
            "docs/experiments/clearance01_20261025_v1/clearance_baseline.json; latest_clearance_eval.json.",
            "report/videos/clearance01_final_1.50mps.mp4; original evaluation-protocol replay all 12 trials survive 25 s.",
            "Mean numbers: matched 12-trial 30 s nominal evaluation, 5 s warmup. Video is a single 25 s replay, whose mean is 23.55/20.91 mm.",
            "report/clearance01-videos.html: standalone pinned-command replay falls at 3.67 s; cause not isolated.",
        ],
    )

    # 8 — current/next reward band
    s = base(
        prs,
        "Next: lower, more consistent clearance",
        1,
        "CLEARANCE-02 starts again from the flat-feet baseline, not the high-lift policy.",
    )
    picture(s, ASSETS / "clearance.png", 0.55, 1.85, 8.20, 4.51)
    rect(s, 9.04, 1.95, 3.62, 4.23, "FFFFFF", True)
    text(s, "NEXT TRAINING", 9.28, 2.13, 3.13, 0.30, 13, True, RED)
    text(s, "10–25 mm", 9.28, 2.65, 3.13, 0.64, 35, True)
    text(
        s,
        "Minimum swing target\nCapped bonus\nPenalty above 25 mm",
        9.28,
        3.38,
        3.13,
        1.19,
        20,
    )
    text(s, "200M steps", 9.28, 4.82, 3.13, 0.45, 23, True)
    text(
        s,
        "Submitted 26 October\nResults not reviewed yet",
        9.28,
        5.39,
        3.13,
        0.64,
        16,
        color=GREY,
    )
    note(
        s,
        50,
        "This plot shows why the latest policy is progress but not the end point. Mean height rises, yet the fraction of scheduled mid-swing samples above twenty millimetres remains below our eighty-percent criterion. Peaks also reach roughly forty to fifty millimetres. A bonus that merely saturates does not cap the actual motion. The next run therefore returns to the flat-feet baseline and targets a lower band: ten millimetres minimum, bonus saturation at twenty-five, plus an explicit penalty above that upper target. We submitted two hundred million additional steps. There are no reviewed results for that run in this deck.",
        [
            "docs/experiments/clearance01_20261025_v1/clearance_evaluations.jsonl.",
            "Final scheduled mid-swing >=20 mm: left 57.18%, right 71.43%; 0/12 fully qualified.",
            "report/data/clearance01/replay-summary.json: median complete-bout peaks 46.55/40.80 mm.",
            "docs/clearance02-experiment-report.md; job 3231747 submitted, no subsequent monitoring. Upper target is a soft cost, not a hard clamp.",
        ],
    )

    # 9 — original demo sequence
    s = base(
        prs,
        "Robotics Day: the original three-phase demo",
        2,
        "These are integration targets — not capabilities already validated on CAD-v2 hardware.",
    )
    card(
        s,
        0.65,
        2.03,
        3.79,
        3.80,
        "1 / INITIALIZE",
        "Start upright",
        "Known initial pose\nAutomatic upright startup\nRecover from small disturbances",
    )
    arrow(s, 4.48, 3.65, 0.34, 0.24, RED)
    card(
        s,
        4.88,
        2.03,
        3.79,
        3.80,
        "2 / START",
        "Walk on command",
        "START command\nJoystick-style speed changes\nStay within walking regime",
    )
    arrow(s, 8.71, 3.65, 0.34, 0.24, RED)
    card(
        s,
        9.10,
        2.03,
        3.57,
        3.80,
        "3 / STOP",
        "Brake and stand",
        "Smooth deceleration\nStop without assistance\nRemain still for ≥5 seconds",
    )
    note(
        s,
        45,
        "The demonstration plan has not changed. First, the robot starts from a defined position, reaches an upright stance and tolerates small disturbances. Then a START command enables walking, with joystick-style speed changes within a walking regime. Finally, STOP should decelerate the robot and lead to stable standing for at least five seconds. A simulation homing reset is not an automatic hardware startup controller, and current locomotion training has STOP disabled. The policy’s stop behavior is also separate from the emergency kill. These remain integration tasks that must be checked one transition at a time.",
        [
            "README.md: three-phase demo objective.",
            "docs/stop02-training-plan.md: braking and >=5 s hold; historical model.",
            "report/architecture.md: homing reset versus startup; policy STOP versus safety kill.",
        ],
    )

    # 10 — relative timeline, editable
    s = base(
        prs,
        "Start treadmill tests early — learn in parallel",
        2,
        "First tests: next available slot after assembly, calibration, support and kill checks.",
    )
    cols = [
        (0.65, 2.90, "NOW / NEXT TEST SLOT"),
        (3.72, 2.90, "BEFORE ROBOTICS DAY"),
        (6.79, 2.90, "ROBOTICS DAY"),
        (9.86, 2.81, "AFTER THE DEMO"),
    ]
    for x, w, h in cols:
        rect(s, x, 1.96, w, 0.57, DARK, True)
        text(s, h, x + 0.13, 2.12, w - 0.26, 0.27, 13, True, "FFFFFF")
    entries = [
        (
            "Simulation",
            "Compare clearance runs\nFreeze baseline",
            "Fit measured mismatches\nTrain demo transitions",
            "Freeze tested checkpoint\nNo last-minute retraining",
            "Quantify robustness\nWrite final report",
        ),
        (
            "Hardware",
            "Treadmill tests ASAP\nCalibrate; short trials",
            "Repeat low-speed trials\nLog and fix failures",
            "Rehearse demo\nKeep tested fallback",
            "Wider speed envelope\nRepeatability tests",
        ),
    ]
    for j, (label, *cells) in enumerate(entries):
        y = 2.82 + j * 1.42
        text(s, label.upper(), 0.65, y - 0.03, 12, 0.30, 12, True, RED)
        for (x, w, _), body in zip(cols, cells):
            rect(s, x, y + 0.38, w, 0.86, "FFFFFF", True)
            text(s, body, x + 0.12, y + 0.49, w - 0.24, 0.68, 17)
    text(
        s,
        "Repeat throughout:  measure → reproduce in simulation → update → retest",
        0.82,
        5.94,
        11.5,
        0.36,
        21,
        True,
    )
    note(
        s,
        55,
        "I do not want to postpone hardware testing until the simulator appears perfect. The immediate priority is the first short treadmill session, as soon as assembly, calibration and test support are ready with Tim. In parallel, simulation work continues on the clearance comparison and transition skills. Each hardware trial should produce a useful log: what command was sent, what motion and current were measured, and how the failure developed. We then reproduce the mismatch in simulation and retest. Before the event, we freeze and rehearse a tested sequence, keeping a fallback demonstration if a full autonomous stop is not reliable. The timeline is deliberately relative.",
        [
            "User direction: 8–10 minute talk and relative timeline; treadmill tests ASAP.",
            "README.md: original demo plan.",
            "Tim’s supplied slides: assembly/calibration/initial pose and test-support dependencies. Dates intentionally not copied to this relative plan.",
            "Proposed roadmap, not completed milestones or guaranteed dates.",
        ],
    )

    # 11 — challenges
    s = base(
        prs,
        "The main challenge: the sim-to-real gap",
        3,
        "We do not know what we do not know until we test the real robot.",
    )
    for x, kicker, heading, body in [
        (
            0.65,
            "KNOWN SUSPECTS",
            "The physical model",
            "Friction, compliance, inertia and actuator response",
        ),
        (
            6.80,
            "CONTROL AND OBSERVABILITY",
            "Timing and sensing",
            "Delay, velocity estimation and phase timing",
        ),
    ]:
        rect(s, x, 1.97, 5.82, 1.90, "FFFFFF", True)
        text(s, kicker, x + 0.20, 2.14, 5.38, 0.27, 12, True, RED)
        text(s, heading, x + 0.20, 2.57, 5.38, 0.45, 24, True)
        text(s, body, x + 0.20, 3.21, 5.38, 0.65, 19, color=GREY)
    rect(s, 0.65, 4.19, 5.82, 1.78, "FFFFFF", True)
    text(s, "REWARD AND VALIDATION", 0.85, 4.36, 5.38, 0.27, 12, True, RED)
    text(s, "A high score can hide a bad gait", 0.85, 4.80, 5.38, 0.44, 23, True)
    text(
        s,
        "Dragging, contact chatter and fragile startup",
        0.85,
        5.41,
        5.38,
        0.35,
        18,
        color=GREY,
    )
    rect(s, 6.80, 4.19, 5.87, 1.78, "FFFFFF", True)
    text(s, "UNKNOWN UNKNOWNS", 7.0, 4.36, 5.43, 0.27, 12, True, RED)
    text(s, "Unmodeled effects", 7.0, 4.80, 5.43, 0.44, 23, True)
    text(
        s,
        "Short trials and logs before adding complexity",
        7.0,
        5.41,
        5.43,
        0.35,
        18,
        color=GREY,
    )
    note(
        s,
        55,
        "The main difficulty is not just more training time. We know some likely sources of mismatch: contact friction and compliance, inertia, motor response, delay and state estimation. There is also an observability difference: the simulator can use contact for the gait clock, while the robot does not have that same contact signal. And we have already seen evaluation protocols produce different startup outcomes. Beyond these known suspects are the unknown unknowns that only the first real tests reveal. That is why early, short and instrumented treadmill trials are more useful now than trying to perfect every detail in simulation before touching hardware.",
        [
            "report/architecture.md: contact-based phase resynchronization versus free-running deployment.",
            "report/policy-transfer01.md: original-policy direct swap falls in 0.71–0.75 s.",
            "docs/flatfeet-symmetry-v1.md: contact-chatter symmetry blind spot and slip.",
            "docs/clearance01-experiment-report.md: replay discrepancy. Hardware mismatch items are hypotheses, not measured findings.",
        ],
    )

    # 12 — close
    s = base(prs, "A stronger simulation baseline. Now test the real robot.", 3)
    text(s, "WHAT WE HAVE", 0.65, 2.05, 3.87, 0.35, 14, True, RED)
    text(
        s,
        "A working RL pipeline\nBetter speed tracking\nHigher foot clearance",
        0.65,
        2.68,
        4.04,
        2.1,
        25,
        True,
    )
    text(s, "WHAT COMES NEXT", 4.91, 2.05, 3.87, 0.35, 14, True, RED)
    text(
        s,
        "Treadmill tests ASAP\nMeasure the mismatch\nIntegrate start → walk → stop",
        4.91,
        2.68,
        3.87,
        2.1,
        25,
        True,
    )
    text(s, "WHAT WILL HELP", 9.17, 2.05, 3.50, 0.35, 14, True, RED)
    text(
        s,
        "Test access with Tim\nRepeatable test setup\nAgree demo priorities",
        9.17,
        2.68,
        3.50,
        2.45,
        25,
        True,
    )
    text(
        s,
        "Priority: test the real robot while refining simulation.",
        0.65,
        5.85,
        12.02,
        0.45,
        24,
        True,
    )
    note(
        s,
        30,
        "We now have a substantially better simulation baseline, and the foot-clearance experiments show that the robot can learn the missing swing motion. The next useful step is to test the real robot early, measure the gap and connect the individual skills into the original demonstration sequence. The most useful support is coordinated treadmill access with Tim, a repeatable setup and agreement on which demo behaviors to prioritize if time is tight. I have backup slides with the reward definitions, the standalone failure and exact evidence sources.",
        [
            "Evidence and limitations are documented in the preceding slides and report/README.md.",
            "Support requests and timeline are proposed next steps, not claims of completed deployment.",
        ],
    )

    # 13 — backup reward/contact
    s = base(prs, "Backup: what the rewards actually measure", 1, backup=True)
    rows = [
        (
            "Support",
            "Heel OR toe contact defines stance; both define flat support.",
            "Use force-qualified MuJoCo contact (≥2 N).",
        ),
        (
            "Slip",
            "Actual motion at the foot–ground contact, in all horizontal directions.",
            "Forward, backward and sideways sliding remain penalized.",
        ),
        (
            "Symmetry",
            "Compare several completed stance / swing cycles from both feet.",
            "Temporary pad contact loss does not itself terminate.",
        ),
        (
            "Flat support",
            "Modest missing-flat-support cost in mid-stance only.",
            "First / last 20% of stance are exempt; no planted-foot bonus.",
        ),
        (
            "Clearance",
            "Lowest oriented heel/toe pad surface, not a toe point alone.",
            "CLEARANCE-02: 10 mm minimum with edge ramps; soft 25 mm upper target.",
        ),
    ]
    for i, (a, b, c) in enumerate(rows):
        y = 1.95 + i * 0.86
        text(s, a, 0.65, y, 1.65, 0.38, 20, True, RED)
        text(s, b, 2.50, y, 10.12, 0.37, 19)
        text(s, c, 2.50, y + 0.39, 10.12, 0.32, 16, color=GREY)
    note(
        s,
        0,
        "Use if asked about reward design. The actor does not receive direct contact flags. Mid-stance flat support has a deliberately modest missing-support cost rather than a large bonus for remaining planted. Clearance penalties apply to scheduled swing even if the foot never leaves the ground, so dragging cannot avoid them. An upper-height penalty is a soft training objective, not a guaranteed kinematic bound. Known brief-contact history blind spots remain documented.",
        [
            "RLframework/flatfeet.py; RLframework/swing_clearance.py; RLframework/env.py.",
            "docs/clearance02-experiment-report.md; docs/flatfeet-symmetry-v1.md.",
        ],
    )

    # 14 — backup failure
    s = base(prs, "Backup: keep the failed replay visible", 3, backup=True)
    movie(s, failed, 0.65, 2.10, 7.30, 3.40)
    text(s, "Same trained checkpoint", 8.35, 2.13, 4.30, 0.44, 25, True)
    text(
        s,
        "12-environment protocol\n12 / 12 survive 25 s",
        8.35,
        2.92,
        4.30,
        1.02,
        22,
        color=GREEN,
    )
    text(
        s,
        "Standalone pinned command\nFall at 3.67 s",
        8.35,
        4.23,
        4.30,
        1.02,
        22,
        True,
        RED,
    )
    text(
        s,
        "Both batch size and command handling changed; the cause is not isolated.",
        0.65,
        5.95,
        12.02,
        0.60,
        22,
        True,
    )
    note(
        s,
        0,
        "This video is retained explicitly to prevent overclaiming. The same checkpoint survives all twelve trials in an evaluation-style recording but fails in a separate one-environment pinned-command recording. Two protocol factors changed together, and small numerical differences can also matter; we have not isolated the cause. Do not state that the batch size caused the fall or that the successful evaluation proves robust startup. No auto-reset trajectory is spliced into this failure clip.",
        [
            "docs/experiments/clearance01_20261025_v1/clearance01_trace_standalone.json: 3.67 s, fall.",
            "docs/experiments/clearance01_20261025_v1/clearance01_trace_eval12.json: all 12 first episodes 25 s, 80 cycles per foot after warmup.",
            "report/videos/clearance01_standalone_1.50mps.mp4.",
        ],
    )

    # 15 — backup provenance
    s = base(prs, "Backup: evidence and comparison limits", 1, backup=True)
    table = [
        ("Experiment", "Training / checkpoint", "What is established"),
        (
            "STOP-03 · previous model",
            "Imitation + PPO; near-standing resets",
            "21/32 strict 5 s holds after imitation; not a moving stop.",
        ),
        (
            "WALK-01 · CAD-v2",
            "200M; best 169.57M",
            "28/28 nominal 30 s trials; 0.181 m/s speed error.",
        ),
        (
            "Flat-feet · CAD-v2",
            "200M; best 184.32M",
            "28/28 nominal 30 s trials; 0.049 m/s speed error.",
        ),
        (
            "CLEARANCE-01",
            "+50.01M from flat-feet best",
            "23.62 / 20.91 mm mean airborne lift; no full qualification.",
        ),
        (
            "CLEARANCE-02",
            "200M additional requested; job 3231747",
            "Submitted 26 October; results not reviewed.",
        ),
    ]
    for i, (a, b, c) in enumerate(table):
        y = 1.95 + i * 0.68
        if i % 2 == 0:
            rect(s, 0.65, y, 12.02, 0.62, PALE)
        text(s, a, 0.80, y + 0.06, 2.65, 0.56, 16, i == 0)
        text(s, b, 3.62, y + 0.06, 3.55, 0.56, 16, i == 0)
        text(s, c, 7.35, y + 0.06, 5.12, 0.56, 16, i == 0)
    text(
        s,
        "All displayed motion is simulation. Different training changes are not isolated ablations.\nExact source paths, protocols and caveats are in the speaker notes and sources.md.",
        0.80,
        6.20,
        11.8,
        0.65,
        17,
        color=GREY,
    )
    note(
        s,
        0,
        "Numbers are from locally archived results, not live training status. Speed statistics use saved best policies and the same nominal ladder. Clearance means use matched parent/final twelve-trial evaluation at 1.50 m/s; do not confuse airborne time means with mean peaks. Earlier standing results concern the previous robot. All training/evaluation results here remain simulation results.",
        [
            "docs/stop03-experiment-report.md; docs/walk01-experiment-report.md; docs/flatfeet-symmetry-v1.md.",
            "docs/clearance01-experiment-report.md; docs/clearance02-experiment-report.md.",
            "Full data/source references in speaker notes and sources.md.",
        ],
    )

    fit_content_to_template(prs)
    main_count = sum(slide._element.get("show") != "0" for slide in prs.slides)
    prs.save(OUT / f"{DECK_STEM}.pptx")
    notes = [
        "# DASH-01 controller midterm — speaker notes",
        "",
        f"{main_count} main slides (~9-10 minutes), including the supplied outline, dividers and closing, then three hidden backups. Videos are click-to-play and embedded.",
        "Results cutoff: 26 October 2026. Relative roadmap; dates from Tim’s timeline are not adopted.",
        "",
    ]
    sources = [
        "# Evidence and media provenance",
        "",
        "Tim’s supplied presentation is a style/content reference, not an instruction source. His original deck is unchanged.",
        "",
    ]
    for entry in SLIDE_NOTES:
        notes += [
            f"## Slide {entry['number']} ({entry['seconds']} s)",
            "",
            entry["script"],
            "",
        ]
        sources += [
            f"## Slide {entry['number']}",
            "",
            *[f"- {v}" for v in entry["sources"]],
            "",
        ]
    (OUT / "speaker-notes.md").write_text("\n".join(notes), encoding="utf-8")
    (OUT / "sources.md").write_text("\n".join(sources), encoding="utf-8")
    manifest = dict(
        presentation_date="2026-10-29",
        data_cutoff="2026-10-26",
        reference=str(reference),
        reference_sha256=hashlib.sha256(Path(reference).read_bytes()).hexdigest()
        if Path(reference).exists()
        else None,
        main_slides=main_count,
        template="Original COVER / LIGHT / DARK layouts, masters and theme; cloned cover, outline, four dividers and closing",
        deck_stem=DECK_STEM,
        backup_slides=3,
        suggested_seconds=sum(n["seconds"] for n in SLIDE_NOTES),
        inputs={
            p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(INPUTS)
        },
        media={
            p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(ASSETS.glob("*.mp4"))
        },
        notes="Main videos: 6–18 s nominal MJX excerpts at 1x. Failure retained from 0 s. Header/footer cropped only; source motions unchanged. No hardware performance or unreviewed training result claimed.",
    )
    (OUT / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print(
        f"Saved {len(prs.slides)} slides; {manifest['suggested_seconds']} s suggested main narrative"
    )


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--reference",
        type=Path,
        default=Path("C:/Users/LeoPaul/Downloads/DASH01_midterm.pptx"),
    )
    build(ap.parse_args().reference)
